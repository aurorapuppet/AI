from typing import List, Dict
from pathlib import Path

# PDF
try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None

# DOCX
try:
    import docx
except Exception:
    docx = None

# PPTX
try:
    from pptx import Presentation
except Exception:
    Presentation = None

# OCR (optional)
try:
    import pytesseract
    from PIL import Image
except Exception:
    pytesseract = None
    Image = None


def extract_text_from_pdf(path: str) -> List[Dict]:
    """Extract text per page from PDF. Returns list of {'page': i, 'text': text}."""
    if PdfReader is None:
        raise RuntimeError("pypdf not installed")
    reader = PdfReader(path)
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ''
        pages.append({'page': i, 'text': text})
    return pages


def extract_text_from_docx(path: str) -> List[Dict]:
    if docx is None:
        raise RuntimeError("python-docx not installed")
    doc = docx.Document(path)
    # group paragraphs; simple approach: keep paragraphs as chunks
    paras = []
    for i, p in enumerate(doc.paragraphs, start=1):
        text = p.text or ''
        paras.append({'page': None, 'text': text})
    return paras


def extract_text_from_pptx(path: str) -> List[Dict]:
    if Presentation is None:
        raise RuntimeError("python-pptx not installed")
    prs = Presentation(path)
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                texts.append(shape.text)
        slides.append({'page': i, 'text': '\n'.join(texts)})
    return slides


def extract_text_generic(path: str) -> List[Dict]:
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == '.pdf':
        return extract_text_from_pdf(path)
    if suffix in ('.docx', '.doc'):
        return extract_text_from_docx(path)
    if suffix in ('.pptx', '.ppt'):
        return extract_text_from_pptx(path)
    # txt fallback
    if suffix == '.txt':
        with open(path, 'r', encoding='utf-8', errors='ignore') as f:
            return [{'page': None, 'text': f.read()}]
    # unknown
    return [{'page': None, 'text': ''}]
